import streamlit as st
from pptx import Presentation
import google.generativeai as genai
import json
import io

# Setup Gemini API securely using Streamlit Secrets
genai.configure(api_key=st.secrets["GEMINI_API_KEY"])

def generate_content(topic, num_slides):
    # Initialize the model
    model = genai.GenerativeModel('gemini-1.5-flash')
    
    # Prompt the AI to return structured JSON
    prompt = f"""
    Create a {num_slides}-slide presentation about: {topic}. 
    Return ONLY a raw JSON list. Do not include markdown formatting like ```json.
    Format the output exactly like this example:
    [
        {{"title": "Introduction to AI", "content": "AI is simulating human intelligence.\\nIt relies on large datasets."}},
        {{"title": "Machine Learning", "content": "A subset of AI.\\nSystems learn from data without explicit programming."}}
    ]
    """
    
    # Get the response
    response = model.generate_content(prompt)
    
    try:
        # Clean the response and parse the JSON string into a Python list
        clean_json = response.text.strip().replace('```json', '').replace('```', '')
        slides_data = json.loads(clean_json)
        return slides_data
    except Exception as e:
        st.error(f"Failed to parse AI output: {e}")
        return None

def create_ppt(title, slides_data):
    prs = Presentation()
    
    # 1. Create Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Generated dynamically with AI"
    
    # 2. Loop through JSON data to create Content Slides
    for slide_dict in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = slide_dict.get("title", "Untitled Slide")
        slide.placeholders[1].text = slide_dict.get("content", "")
        
    # 3. Save to memory buffer
    ppt_stream = io.BytesIO()
    prs.save(ppt_stream)
    ppt_stream.seek(0)
    return ppt_stream

# --- User Interface ---
st.title("🤖 AI-Powered PPT Generator")
st.write("Type a topic, and AI will write the slides and generate the PowerPoint for you.")

topic = st.text_input("What is your presentation about?", "The impact of renewable energy on the economy")
num_slides = st.slider("Number of content slides", min_value=3, max_value=10, value=5)

if st.button("Generate AI Presentation"):
    with st.spinner("AI is researching and writing your slides..."):
        
        # Step A: Get content from Gemini
        slides_data = generate_content(topic, num_slides)
        
        if slides_data:
            # Step B: Build the PPT file
            ppt_file = create_ppt(topic, slides_data)
            
            # Step C: Offer the download
            st.success("Your presentation is ready!")
            st.download_button(
                label="⬇️ Download .pptx File",
                data=ppt_file,
                file_name=f"{topic.replace(' ', '_')}_deck.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
